#!/usr/bin/env python3

from pathlib import Path
from openpyxl import load_workbook

ROOT = Path("data/raw/crispr_datasets")

for path in sorted(ROOT.rglob("*.xlsx")):

    print("\n" + "=" * 100)
    print(path)
    print("=" * 100)

    try:
        wb = load_workbook(
            path,
            read_only=True,
            data_only=True
        )

        print("Sheets:", len(wb.sheetnames))

        for ws in wb.worksheets:

            print(
                f"\n  SHEET: {ws.title}"
                f" | state={ws.sheet_state}"
                f" | rows={ws.max_row}"
                f" | cols={ws.max_column}"
            )

            # Primeras 8 filas
            for i, row in enumerate(
                ws.iter_rows(min_row=1, max_row=min(8, ws.max_row), values_only=True),
                start=1
            ):
                values = [
                    str(x)[:150] if x is not None else ""
                    for x in row[:min(ws.max_column, 20)]
                ]
                print(f"    row {i}: {' | '.join(values)}")

        wb.close()

    except Exception as e:
        print("ERROR:", repr(e))
