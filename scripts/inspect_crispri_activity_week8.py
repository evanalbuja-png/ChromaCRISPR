#!/usr/bin/env python3

from pathlib import Path
from openpyxl import load_workbook
import pandas as pd
import re

ROOT = Path("data/raw/crispr_datasets")
MATRIX = Path("data/processed/sgRNA_feature_matrix_phase1.csv")
LOG = Path("logs/week8_log.md")

TARGET = "CRISPRi activity score"

print("=" * 100)
print("STEP 1 — IDENTIFICACIÓN E INSPECCIÓN DE CRISPRi ACTIVITY SCORE")
print("=" * 100)

# ---------------------------------------------------------------------
# 1. Identify exact Excel file containing the target column
# ---------------------------------------------------------------------
matches = []

for path in sorted(ROOT.rglob("*.xlsx")):
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            headers = [
                str(cell.value).strip() if cell.value is not None else ""
                for cell in next(ws.iter_rows(min_row=1, max_row=1))
            ]

            for header in headers:
                if TARGET.lower() in header.lower():
                    matches.append((path, ws.title, headers))
        wb.close()
    except Exception as e:
        print(f"WARNING: could not inspect {path}: {e}")

print("\n[1] ARCHIVOS CON COLUMNA CRISPRi ACTIVITY SCORE")
print("-" * 100)

if not matches:
    raise SystemExit("ERROR: No se encontró ninguna columna CRISPRi activity score.")

for path, sheet, headers in matches:
    print(f"FILE : {path}")
    print(f"SHEET: {sheet}")
    print(f"COLS : {headers}")
    print()

if len(matches) != 1:
    raise SystemExit(
        f"ERROR: Se encontraron {len(matches)} coincidencias. "
        "No se continuará hasta resolver la ambigüedad."
    )

activity_file, activity_sheet, activity_headers = matches[0]

print(f"Archivo exacto identificado: {activity_file}")
print(f"Hoja exacta: {activity_sheet}")

# ---------------------------------------------------------------------
# 2. Load activity-score sheet and display columns + first rows
# ---------------------------------------------------------------------
print("\n[2] COLUMNAS Y PRIMERAS FILAS")
print("-" * 100)

wb = load_workbook(activity_file, read_only=True, data_only=True)
ws = wb[activity_sheet]

rows = ws.iter_rows(values_only=True)
headers = list(next(rows))

records = []
for row in rows:
    records.append(row)

activity_df = pd.DataFrame(records, columns=headers)
wb.close()

print(f"Shape: {activity_df.shape}")
print("\nColumnas:")
for i, col in enumerate(activity_df.columns, 1):
    print(f"{i:2d}. {col}")

print("\nPrimeras 10 filas:")
print(activity_df.head(10).to_string(index=False))

# ---------------------------------------------------------------------
# 3. Load principal matrix
# ---------------------------------------------------------------------
print("\n[3] MATRIZ PRINCIPAL")
print("-" * 100)

matrix = pd.read_csv(MATRIX, low_memory=False)

print(f"Archivo: {MATRIX}")
print(f"Shape: {matrix.shape}")

print("\nColumnas:")
for i, col in enumerate(matrix.columns, 1):
    print(f"{i:2d}. {col}")

# ---------------------------------------------------------------------
# 4. Identify sequence / gene columns robustly
# ---------------------------------------------------------------------
def find_column(df, candidates):
    normalized = {
        re.sub(r"[^a-z0-9]", "", str(c).lower()): c
        for c in df.columns
    }

    for candidate in candidates:
        key = re.sub(r"[^a-z0-9]", "", candidate.lower())
        if key in normalized:
            return normalized[key]

    return None


activity_seq = find_column(
    activity_df,
    [
        "sgRNA sequence",
        "guide_sequence",
        "guide sequence",
        "sgRNA_sequence",
        "spacer",
        "Spacer",
    ],
)

activity_gene = find_column(
    activity_df,
    [
        "gene symbol",
        "gene_symbol",
        "gene",
        "target_gene",
    ],
)

matrix_seq = find_column(
    matrix,
    [
        "guide_sequence",
        "guide sequence",
        "sgRNA sequence",
        "sgRNA_sequence",
        "spacer",
        "Spacer",
    ],
)

matrix_gene = find_column(
    matrix,
    [
        "gene_symbol",
        "gene symbol",
        "gene",
        "target_gene",
    ],
)

print("\nColumnas detectadas para unión:")
print(f"Activity sequence : {activity_seq}")
print(f"Activity gene     : {activity_gene}")
print(f"Matrix sequence   : {matrix_seq}")
print(f"Matrix gene       : {matrix_gene}")

if activity_seq is None or matrix_seq is None:
    raise SystemExit(
        "ERROR: No fue posible identificar columnas de secuencia en ambos datasets."
    )

# ---------------------------------------------------------------------
# 5. Normalize sequences for exact comparison
# ---------------------------------------------------------------------
def normalize_sequence(series):
    return (
        series.astype("string")
        .str.strip()
        .str.upper()
        .str.replace(r"[^ACGTN]", "", regex=True)
    )


activity_df["_seq_norm"] = normalize_sequence(activity_df[activity_seq])
matrix["_seq_norm"] = normalize_sequence(matrix[matrix_seq])

# ---------------------------------------------------------------------
# 6. Number of gRNAs
# ---------------------------------------------------------------------
activity_unique = activity_df["_seq_norm"].dropna()
activity_unique = activity_unique[activity_unique != ""].nunique()

matrix_unique = matrix["_seq_norm"].dropna()
matrix_unique = matrix_unique[matrix_unique != ""].nunique()

print("\n[4] NÚMERO DE gRNAs")
print("-" * 100)
print(f"Activity-score rows             : {len(activity_df):,}")
print(f"Activity-score unique sequences : {activity_unique:,}")
print(f"Matrix unique sequences         : {matrix_unique:,}")

# ---------------------------------------------------------------------
# 7. Potential matches by sequence
# ---------------------------------------------------------------------
activity_seq_set = set(
    activity_df.loc[
        activity_df["_seq_norm"].notna() & (activity_df["_seq_norm"] != ""),
        "_seq_norm"
    ]
)

matrix_seq_set = set(
    matrix.loc[
        matrix["_seq_norm"].notna() & (matrix["_seq_norm"] != ""),
        "_seq_norm"
    ]
)

sequence_overlap = activity_seq_set & matrix_seq_set

print("\n[5] MATCH POR guide_sequence")
print("-" * 100)
print(f"Activity gRNAs matching matrix: {len(sequence_overlap):,}")
print(f"Percent of activity dataset    : "
      f"{100 * len(sequence_overlap) / max(activity_unique, 1):.2f}%")
print(f"Percent of matrix              : "
      f"{100 * len(sequence_overlap) / max(matrix_unique, 1):.2f}%")

# ---------------------------------------------------------------------
# 8. Gene-level / sequence+gene potential matching
# ---------------------------------------------------------------------
if activity_gene and matrix_gene:
    activity_df["_gene_norm"] = (
        activity_df[activity_gene]
        .astype("string")
        .str.strip()
        .str.upper()
    )

    matrix["_gene_norm"] = (
        matrix[matrix_gene]
        .astype("string")
        .str.strip()
        .str.upper()
    )

    activity_pairs = set(
        zip(
            activity_df["_seq_norm"].fillna(""),
            activity_df["_gene_norm"].fillna("")
        )
    )

    matrix_pairs = set(
        zip(
            matrix["_seq_norm"].fillna(""),
            matrix["_gene_norm"].fillna("")
        )
    )

    pair_overlap = activity_pairs & matrix_pairs

    print("\n[6] MATCH POR guide_sequence + gene_symbol")
    print("-" * 100)
    print(f"Matching sequence+gene pairs: {len(pair_overlap):,}")
    print(
        f"Percent of activity dataset: "
        f"{100 * len(pair_overlap) / max(len(activity_pairs), 1):.2f}%"
    )

else:
    pair_overlap = set()
    print("\n[6] MATCH POR guide_sequence + gene_symbol")
    print("-" * 100)
    print("No fue posible evaluar combinación sequence + gene.")

# ---------------------------------------------------------------------
# 9. Activity score distribution
# ---------------------------------------------------------------------
score_col = find_column(
    activity_df,
    [
        "CRISPRi activity score",
        "CRISPRi activity score [Horlbeck et al., eLife 2016]",
    ],
)

if score_col is None:
    raise SystemExit("ERROR: No se pudo localizar la columna de activity score.")

scores = pd.to_numeric(activity_df[score_col], errors="coerce").dropna()

print("\n[7] EVALUACIÓN DEL ACTIVITY SCORE")
print("-" * 100)
print(f"Score column : {score_col}")
print(f"Numeric scores: {len(scores):,}")
print(f"Missing/non-numeric: {len(activity_df) - len(scores):,}")
print(f"Unique scores: {scores.nunique():,}")
print(f"Minimum: {scores.min()}")
print(f"Maximum: {scores.max()}")
print(f"Mean: {scores.mean()}")
print(f"Median: {scores.median()}")
print(f"Std: {scores.std()}")
print(f"Q01: {scores.quantile(0.01)}")
print(f"Q25: {scores.quantile(0.25)}")
print(f"Q75: {scores.quantile(0.75)}")
print(f"Q99: {scores.quantile(0.99)}")

if scores.min() >= 0 and scores.max() <= 1:
    normalization_assessment = (
        "El score está numéricamente acotado entre 0 y 1. "
        "Esto es consistente con un score normalizado/bounded, "
        "pero esta inspección por sí sola no demuestra el método de normalización."
    )
else:
    normalization_assessment = (
        "El score no está acotado entre 0 y 1; parece ser un score raw/no "
        "normalizado a [0,1]. La confirmación definitiva requiere la documentación "
        "del dataset original."
    )

print("\nEvaluación:")
print(normalization_assessment)

# ---------------------------------------------------------------------
# 10. Save concise inspection report
# ---------------------------------------------------------------------
REPORT = Path("results/week8_crispri_activity_inspection.txt")
REPORT.parent.mkdir(parents=True, exist_ok=True)

with REPORT.open("w") as f:
    f.write("ChromaCRISPR Phase 1 — Week 8 CRISPRi activity inspection\n")
    f.write("=" * 80 + "\n")
    f.write(f"Activity file: {activity_file}\n")
    f.write(f"Activity sheet: {activity_sheet}\n")
    f.write(f"Activity rows: {len(activity_df)}\n")
    f.write(f"Unique activity sequences: {activity_unique}\n")
    f.write(f"Matrix unique sequences: {matrix_unique}\n")
    f.write(f"Sequence overlap: {len(sequence_overlap)}\n")
    f.write(f"Sequence+gene overlap: {len(pair_overlap)}\n")
    f.write(f"Score column: {score_col}\n")
    f.write(f"Score min: {scores.min()}\n")
    f.write(f"Score max: {scores.max()}\n")
    f.write(f"Score mean: {scores.mean()}\n")
    f.write(f"Score median: {scores.median()}\n")
    f.write(f"Score assessment: {normalization_assessment}\n")

# ---------------------------------------------------------------------
# 11. Update Week 8 log
# ---------------------------------------------------------------------
LOG.parent.mkdir(parents=True, exist_ok=True)

with LOG.open("a") as f:
    f.write("\n## Week 8 — CRISPRi activity score inspection\n\n")
    f.write(f"- Exact file: `{activity_file}`\n")
    f.write(f"- Sheet: `{activity_sheet}`\n")
    f.write(f"- Activity-score rows: {len(activity_df):,}\n")
    f.write(f"- Unique activity gRNA sequences: {activity_unique:,}\n")
    f.write(f"- Matrix unique guide sequences: {matrix_unique:,}\n")
    f.write(f"- Potential exact sequence matches: {len(sequence_overlap):,}\n")
    f.write(f"- Potential sequence + gene matches: {len(pair_overlap):,}\n")
    f.write(f"- Activity score column: `{score_col}`\n")
    f.write(f"- Score minimum: {scores.min()}\n")
    f.write(f"- Score maximum: {scores.max()}\n")
    f.write(f"- Score mean: {scores.mean()}\n")
    f.write(f"- Score median: {scores.median()}\n")
    f.write(f"- Normalization assessment: {normalization_assessment}\n")

print("\n" + "=" * 100)
print("Inspección CRISPRi activity score completada")
print("=" * 100)
